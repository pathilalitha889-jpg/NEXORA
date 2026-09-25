
import sqlite3
import os
import smtplib
from email.message import EmailMessage
from docx import Document
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from dotenv import load_dotenv
from openai import OpenAI
from datetime import datetime

load_dotenv()

client = OpenAI()

EMAIL_ADDRESS = os.getenv("EMAIL_ADDRESS")
EMAIL_APP_PASSWORD = os.getenv("EMAIL_APP_PASSWORD")

SCOPES = ["https://www.googleapis.com/auth/calendar.events"]


def get_calendar_service():
    creds = None

    if os.path.exists("token.json"):
        creds = Credentials.from_authorized_user_file(
            "token.json",
            SCOPES
        )

    if not creds or not creds.valid:

        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())

        else:
            flow = InstalledAppFlow.from_client_secrets_file(
                "credentials.json",
                SCOPES
            )

            creds = flow.run_local_server(port=0)

        with open("token.json", "w") as token:
            token.write(creds.to_json())

    return build(
        "calendar",
        "v3",
        credentials=creds
    )


conn = sqlite3.connect("nexora.db")
cursor = conn.cursor()


cursor.execute("""
CREATE TABLE IF NOT EXISTS meetings (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    person TEXT,
    date TEXT,
    time TEXT
)
""")


cursor.execute("""
CREATE TABLE IF NOT EXISTS tasks (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    task TEXT
)
""")

cursor.execute("""
CREATE TABLE IF NOT EXISTS memories (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    information TEXT
)
""")

conn.commit()



def ask_permission(action):
    print(f"\nNEXORA wants to perform: {action}")
    choice = input("Do you want to continue? (yes/no): ")

    return choice.lower() == "yes"


def add_task(task):
    cursor.execute(
        "INSERT INTO tasks (task) VALUES (?)",
        (task,)
    )
    conn.commit()

    print("\nTask Added!")
    print("Task:", task)


def save_memory(information):
    cursor.execute(
        "INSERT INTO memories (information) VALUES (?)",
        (information,)
    )
    conn.commit()

    print("Memory saved!")
def get_memories():
    cursor.execute("SELECT information FROM memories")
    memories = cursor.fetchall()

    return [memory[0] for memory in memories]
def read_document():

    document = Document("questions.docx")

    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text


def search_document(question):

    document_text = read_document()

    question_words = (
        question.lower()
        .replace("?", "")
        .replace(".", "")
        .replace(",", "")
        .split()
    )

    words = set(question_words)

    best_score = 0
    best_text = ""

    paragraphs = document_text.split("\n")

    for paragraph in paragraphs:

        paragraph_lower = paragraph.lower()

        score = 0

        for word in words:

            if len(word) > 2 and word in paragraph_lower:
                score += 1

        if score > best_score:
            best_score = score
            best_text = paragraph

    if best_score == 0:
        return document_text

    return best_text

def schedule_meeting(person, date, time):

    cursor.execute(
        "INSERT INTO meetings (person, date, time) VALUES (?, ?, ?)",
        (person, date, time)
    )

    conn.commit()

    service = get_calendar_service()

    event = {
        "summary": f"Meeting with {person}",
        "description": "Meeting scheduled by NEXORA AI Agent",

        "start": {
            "dateTime": f"{date}T{time}:00+05:30",
            "timeZone": "Asia/Kolkata"
        },

        "end": {
            "dateTime": f"{date}T{time}:00+05:30",
            "timeZone": "Asia/Kolkata"
        }
    }

    created_event = service.events().insert(
        calendarId="primary",
        body=event
    ).execute()

    print("\nMeeting Scheduled!")
    print("Person:", person)
    print("Date:", date)
    print("Time:", time)
    print("Google Calendar Event Created!")
    print("Calendar Link:", created_event.get("htmlLink"))


def show_tasks():

    cursor.execute("SELECT * FROM tasks")

    tasks = cursor.fetchall()

    print("\nYour Tasks:")

    for task in tasks:
        print(task)


def show_meetings():

    cursor.execute("SELECT * FROM meetings")

    meetings = cursor.fetchall()

    print("\nYour Meetings:")

    for meeting in meetings:
        print(meeting)


def send_email(recipient, message):

    email = EmailMessage()

    email["From"] = EMAIL_ADDRESS
    email["To"] = recipient
    email["Subject"] = "Message from NEXORA"

    email.set_content(message)

    with smtplib.SMTP_SSL(
        "smtp.gmail.com",
        465
    ) as server:

        server.login(
            EMAIL_ADDRESS,
            EMAIL_APP_PASSWORD
        )

        server.send_message(email)

    print("\nEmail Sent Successfully!")


print("Welcome to NEXORA")


name = input("Enter your name: ")

command = input("What do you want me to do? ")


response = client.responses.create(

    model="gpt-5.6-luna",

    input=f"""
You are NEXORA, an intelligent AI assistant.

User request:
{command}

Classify it as exactly one of:

MEETING
ADD_TASK
SHOW_MEETING
SHOW_TASK
SEND_EMAIL
MEMORY
RAG
OTHER

Use RAG when the user asks a question about information that may be present in the uploaded document "questions.docx".

Examples of RAG:
- What is JDBC architecture?
- Explain JDBC drivers.
- What are the types of JDBC drivers?
- Explain JDBC connection steps.
- Give the answer for JDBC question.
- What is the difference between Statement and PreparedStatement?

RAG means the user wants an answer from the uploaded document.


Return only the classification.
"""
)

action = response.output_text.strip()

print("\nAI understood:", action)

if action == "MEMORY":

    memory_type = client.responses.create(
        model="gpt-5.6-luna",
        input=f"""
Determine whether the user wants to SAVE a new memory or RETRIEVE existing memories.

User request:
{command}

Return exactly one:
SAVE
RETRIEVE
"""
    )

    memory_action = memory_type.output_text.strip()

    if memory_action == "SAVE":

        memory_response = client.responses.create(
            model="gpt-5.6-luna",
            input=f"""
Extract the important information the user wants NEXORA to remember.

User request:
{command}

Return only the information to remember.
"""
        )

        information = memory_response.output_text.strip()

        save_memory(information)

    elif memory_action == "RETRIEVE":

        memories = get_memories()

        if memories:
            print("\nNEXORA Memory:")

            for memory in memories:
                print("-", memory)

        else:
            print("No memories found.")


elif action == "RAG":

    relevant_text = search_document(command)

    rag_response = client.responses.create(

        model="gpt-5.6-luna",

        input=f"""
You are NEXORA.

Answer the user's question using the document content below.

DOCUMENT CONTENT:
{relevant_text}

USER QUESTION:
{command}

Give a clear and accurate answer.

If the information is not available, say:
Sorry, I could not find that information in the document.
"""
    )

    print("\nNEXORA Answer:")
    print(rag_response.output_text)
    
elif action == "ADD_TASK":

    task_response = client.responses.create(
        model="gpt-5.6-luna",
        input=f"""
Extract the task from this request:

{command}

Return only the task text.
"""
    )

    task = task_response.output_text.strip()

    add_task(task)



elif action == "SEND_EMAIL":

    email_response = client.responses.create(
        model="gpt-5.6-luna",
        input=f"""
Extract the email details from this request:

{command}

If the user says "your email address" or "my email address",
use the configured email address.

Return exactly:

RECIPIENT: name or email
MESSAGE: message
"""
    )

    lines = email_response.output_text.strip().split("\n")

    recipient = lines[0].replace(
        "RECIPIENT:",
        ""
    ).strip()

    if recipient.lower() in ["your email address", "my email address"]:
        recipient = EMAIL_ADDRESS

    message = lines[1].replace(
        "MESSAGE:",
        ""
    ).strip()

    if ask_permission("Send email"):
        send_email(recipient, message)
    else:
        print("Email cancelled.")

elif action == "MEETING":

    details = client.responses.create(

        model="gpt-5.6-luna",

        input=f"""
Today is {datetime.now().strftime('%Y-%m-%d')}.

Extract the meeting details from this request:

{command}

Convert relative dates like tomorrow, today, next Monday into YYYY-MM-DD.

Convert time into 24-hour HH:MM format.

Return exactly:

PERSON: name
DATE: YYYY-MM-DD
TIME: HH:MM
"""
    )

    lines = details.output_text.strip().split("\n")

    person = lines[0].replace(
        "PERSON:",
        ""
    ).strip()

    date = lines[1].replace(
        "DATE:",
        ""
    ).strip()

    time = lines[2].replace(
        "TIME:",
        ""
    ).strip()

    if ask_permission("Create Google Calendar meeting"):

        schedule_meeting(
            person,
            date,
            time
        )

    else:

        print("Meeting cancelled.")


elif action == "SHOW_TASK":

    show_tasks()


elif action == "SHOW_MEETING":

    show_meetings()


else:

    print("Sorry, I don't know that task yet.")

conn.close()
