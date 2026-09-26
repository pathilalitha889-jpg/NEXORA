import os
from datetime import datetime
from typing import TypedDict, Optional

from dotenv import load_dotenv
from openai import OpenAI
from langgraph.graph import StateGraph, START, END
from pypdf import PdfReader
from docx import Document

import tools


# =========================================================
# CONFIGURATION
# =========================================================

load_dotenv()

try:
    client = OpenAI()
except Exception:
    client = None


VALID_ACTIONS = {
    "MEETING",
    "ADD_TASK",
    "SHOW_MEETING",
    "SHOW_TASK",
    "SEND_EMAIL",
    "MEMORY",
    "RAG",
    "OTHER"
}


# =========================================================
# AGENT STATE
# =========================================================

class AgentState(TypedDict, total=False):
    command: str
    action: str
    result: str
    file_path: Optional[str]
    permission_granted: bool
    google_access_token: Optional[str]


# =========================================================
# HELPERS
# =========================================================

def use_ai() -> bool:
    return client is not None


def normalize_action(text: str) -> str:

    text = text.strip().upper()

    if text in VALID_ACTIONS:
        return text

    for action in VALID_ACTIONS:

        if action in text:
            return action

    return "OTHER"


# =========================================================
# LOCAL FALLBACK CLASSIFIER
# =========================================================

def local_classify(command: str) -> str:

    text = command.lower().strip()

    # -----------------------------------------------------
    # MEETING
    # -----------------------------------------------------

    meeting_words = [
        "meeting",
        "appointment",
        "calendar",
        "schedule",
        "scheduled",
        "book",
        "arrange",
        "organize",
        "meet"
    ]

    if any(word in text for word in meeting_words):
        return "MEETING"

    # -----------------------------------------------------
    # SHOW TASK
    # -----------------------------------------------------

    show_task_patterns = [
        "what are my tasks",
        "what do i need to do",
        "show my tasks",
        "show tasks",
        "list my tasks",
        "my task list",
        "tasks do i have",
        "tasks that i have",
        "saved tasks",
        "pending tasks",
        "to do list",
        "todo list"
    ]

    if any(pattern in text for pattern in show_task_patterns):
        return "SHOW_TASK"

    # -----------------------------------------------------
    # SHOW MEETING
    # -----------------------------------------------------

    show_meeting_patterns = [
        "show my meetings",
        "show meetings",
        "list my meetings",
        "what meetings do i have",
        "my meetings",
        "scheduled meetings",
        "what is scheduled",
        "anything scheduled"
    ]

    if any(pattern in text for pattern in show_meeting_patterns):
        return "SHOW_MEETING"

    # -----------------------------------------------------
    # SEND EMAIL
    # -----------------------------------------------------

    email_words = [
        "send an email",
        "send email",
        "email",
        "mail",
        "send a message",
        "send message"
    ]

    if any(word in text for word in email_words):
        return "SEND_EMAIL"

    # -----------------------------------------------------
    # MEMORY
    # -----------------------------------------------------

    memory_words = [
        "remember",
        "memorize",
        "keep in mind",
        "save this",
        "store this",
        "what do you remember",
        "do you remember",
        "what have i told you"
    ]

    if any(word in text for word in memory_words):
        return "MEMORY"

    # -----------------------------------------------------
    # RAG / DOCUMENT
    # -----------------------------------------------------

    rag_words = [
        "pdf",
        "document",
        "uploaded file",
        "uploaded document",
        "docx",
        "jdbc",
        "batch updates",
        "from my document",
        "from the document",
        "from the pdf"
    ]

    if any(word in text for word in rag_words):
        return "RAG"

    # -----------------------------------------------------
    # ADD TASK
    # -----------------------------------------------------

    task_words = [
        "task",
        "todo",
        "to-do",
        "remind me",
        "i need to",
        "i have to",
        "i want to",
        "i should",
        "don't let me forget",
        "need to",
        "have to",
        "should",
        "must",
        "prepare",
        "study",
        "learn",
        "practice",
        "revise",
        "finish",
        "complete",
        "submit",
        "work on"
    ]

    if any(word in text for word in task_words):

        # Avoid treating questions about existing tasks as ADD_TASK
        if any(word in text for word in [
            "what are",
            "what do i need",
            "show",
            "list",
            "saved",
            "pending"
        ]):
            return "SHOW_TASK"

        return "ADD_TASK"

    return "OTHER"


# =========================================================
# PERMISSION
# =========================================================

def ask_permission(
    action: str,
    state: AgentState
) -> bool:

    # Web UI explicitly grants permission
    if state.get("permission_granted", False):
        return True

    # Terminal mode
    if os.getenv("NEXORA_UI") != "1":

        print(
            f"\nPermission required for: {action}"
        )

        choice = input(
            "Do you want to allow this action? (yes/no): "
        )

        return choice.lower().strip() == "yes"

    return False


# =========================================================
# ANALYZE REQUEST
# =========================================================

def analyze_request(state):

    command = state.get(
        "command",
        ""
    ).strip()

    if not command:

        return {
            "action": "OTHER",
            "result": "Please enter a command."
        }

    print("\nAnalyzing request...")

    prompt = f"""
You are NEXORA's intelligent intent router.

Your job is to understand the user's MEANING
and select the correct tool.

IMPORTANT:
- Do NOT depend on exact keywords.
- Understand natural language.
- Understand paraphrasing.
- Understand short requests.
- Understand indirect requests.
- Understand different sentence structures.
- Understand requests such as:
  "I want to study Java tomorrow"
  "I need to prepare for my interview"
  "Please put this on my task list"
  "Don't let me forget my assignment"
- Focus on what the user wants NEXORA to DO.
- Do not return OTHER when a supported action clearly matches.

Choose exactly ONE:

MEETING
ADD_TASK
SHOW_MEETING
SHOW_TASK
SEND_EMAIL
MEMORY
RAG
OTHER

MEETING:
Schedule, arrange, book, create, organize, or plan
a meeting, appointment, or calendar event.

ADD_TASK:
The user wants to add something to a task list
or create something that they need to do.

Examples:
"I want to study Java tomorrow"
"I need to complete my project"
"Make sure I prepare for my interview"
"Don't let me forget my assignment"
"Please add studying DSA to my tasks"

SHOW_TASK:
The user wants to see existing tasks.

Examples:
"What do I need to do?"
"What are my tasks?"
"Show my saved tasks"
"List everything on my task list"

SHOW_MEETING:
The user wants to see existing meetings.

Examples:
"What meetings do I have?"
"Anything scheduled?"
"Show my appointments"

SEND_EMAIL:
The user wants NEXORA to send an email or message.

Examples:
"Tell Ravi that I will be late"
"Please send an email to Ravi"
"Mail this person"

MEMORY:
The user wants NEXORA to save or retrieve memory.

Examples:
"Remember that I am learning Java"
"Keep this in mind"
"What do you remember about me?"
"Do you know what I am learning?"

RAG:
The user wants an answer from the uploaded document/PDF.

Examples:
"What does the PDF say about JDBC?"
"Explain batch updates from my document"
"Answer this from the uploaded document"

OTHER:
Only when the request does not match any supported tool.

USER REQUEST:
{command}

Return ONLY the action name.
"""

    # -----------------------------------------------------
    # AI CLASSIFICATION
    # -----------------------------------------------------

    if use_ai():

        try:

            response = client.responses.create(
                model="gpt-5.6-luna",
                input=prompt
            )

            raw_action = (
                response.output_text
                .strip()
                .upper()
            )

            action = normalize_action(
                raw_action
            )

            # -------------------------------------------------
            # If AI returns OTHER, check local fallback too
            # -------------------------------------------------

            if action == "OTHER":

                fallback_action = local_classify(
                    command
                )

                if fallback_action != "OTHER":
                    action = fallback_action

            print(
                "NEXORA understood:",
                action
            )

            return {
                "action": action
            }

        except Exception as e:

            print(
                "AI classification error:",
                e
            )

    # -----------------------------------------------------
    # LOCAL FALLBACK
    # -----------------------------------------------------

    fallback_action = local_classify(
        command
    )

    print(
        "NEXORA local fallback:",
        fallback_action
    )

    return {
        "action": fallback_action
    }


# =========================================================
# MEETING NODE
# =========================================================

def meeting_node(state: AgentState):

    command = state.get(
        "command",
        ""
    ).strip()

    person = ""
    date = ""
    time = ""

    # -----------------------------------------
    # AI extraction
    # -----------------------------------------

    if use_ai():

        try:

            details = client.responses.create(
                model="gpt-5.6-luna",
                input=f"""
Today is {datetime.now().strftime('%Y-%m-%d')}.

Extract the meeting details from this request:

{command}

Convert relative dates such as tomorrow, today
and next Monday into YYYY-MM-DD.

Convert time into 24-hour HH:MM format.

Return exactly:

PERSON: name
DATE: YYYY-MM-DD
TIME: HH:MM
"""
            )

            lines = (
                details.output_text
                .strip()
                .splitlines()
            )

            for line in lines:

                upper_line = line.upper()

                if upper_line.startswith("PERSON:"):

                    person = (
                        line.split(
                            ":",
                            1
                        )[1].strip()
                    )

                elif upper_line.startswith("DATE:"):

                    date = (
                        line.split(
                            ":",
                            1
                        )[1].strip()
                    )

                elif upper_line.startswith("TIME:"):

                    time = (
                        line.split(
                            ":",
                            1
                        )[1].strip()
                    )

        except Exception as e:

            print(
                "Meeting AI extraction failed:",
                e
            )

    # -----------------------------------------
    # Basic fallback extraction
    # -----------------------------------------

    if not person:

        words = command.split()

        lowered_words = [
            word.lower()
            for word in words
        ]

        if "with" in lowered_words:

            try:

                index = lowered_words.index(
                    "with"
                )

                if index + 1 < len(words):

                    person = words[
                        index + 1
                    ]

            except Exception:
                pass

    if not person:
        person = "Unknown"

    if not date:
        date = datetime.now().strftime(
            "%Y-%m-%d"
        )

    if not time:
        time = "10:00"

    print("\nMeeting Details:")
    print("Person:", person)
    print("Date:", date)
    print("Time:", time)

    # -----------------------------------------
    # Permission
    # -----------------------------------------

    allowed = ask_permission(
        f"Create a Google Calendar meeting with {person}",
        state
    )

    if not allowed:

        if os.getenv("NEXORA_UI") == "1":

            return {
                "result": "PERMISSION_REQUIRED"
            }

        return {
            "result": "Meeting cancelled by user."
        }

    print("Permission granted.")

    try:

        tools.save_meeting(
            person,
            date,
            time
        )

        link = tools.create_calendar_event(
            person,
            date,
            time,
            state.get(
                "google_access_token"
            )
        )

        print(
            "\nGoogle Calendar Event Created!"
        )

        print(
            "Calendar Link:",
            link
        )

        return {
            "result": link
        }

    except Exception as e:

        print(
            "Calendar error:",
            e
        )

        return {
            "result": f"Calendar error: {e}"
        }


# =========================================================
# ADD TASK NODE
# =========================================================

def task_node(state: AgentState):

    command = state.get(
        "command",
        ""
    ).strip()

    task = command

    if use_ai():

        try:

            task_response = client.responses.create(
                model="gpt-5.6-luna",
                input=f"""
Extract only the actual task from this request.

User request:
{command}

Remove phrases such as:
"I want to",
"I need to",
"please",
"add a task to",
"don't let me forget to"

Return only the task text.
"""
            )

            extracted_task = (
                task_response.output_text
                .strip()
            )

            if extracted_task:
                task = extracted_task

        except Exception as e:

            print(
                "Task AI extraction failed:",
                e
            )

    if not task:
        return {
            "result": "Task description is missing."
        }

    try:

        tools.add_task(task)

    except Exception as e:

        print(
            "Task save error:",
            e
        )

        return {
            "result": "Unable to add the task."
        }

    print("\nTask Added!")
    print("Task:", task)

    return {
        "result": f"Task Added: {task}"
    }


# =========================================================
# SHOW TASK NODE
# =========================================================

def show_task_node(state: AgentState):

    try:

        tasks = tools.get_tasks()

    except Exception as e:

        print(
            "Task retrieval error:",
            e
        )

        return {
            "result": "Unable to retrieve tasks."
        }

    print("\nYour Tasks:")

    if tasks:

        for task in tasks:
            print(task)

    else:

        print("No tasks found.")

    return {
        "result": str(tasks)
    }


# =========================================================
# SHOW MEETING NODE
# =========================================================

def show_meeting_node(state: AgentState):

    try:

        meetings = tools.get_meetings()

    except Exception as e:

        print(
            "Meeting retrieval error:",
            e
        )

        return {
            "result": "Unable to retrieve meetings."
        }

    print("\nYour Meetings:")

    if meetings:

        for meeting in meetings:
            print(meeting)

    else:

        print("No meetings found.")

    return {
        "result": str(meetings)
    }


# =========================================================
# MEMORY NODE
# =========================================================

def memory_node(state: AgentState):

    command = state.get(
        "command",
        ""
    ).strip()

    memory_action = ""

    # -----------------------------------------
    # AI classification
    # -----------------------------------------

    if use_ai():

        try:

            memory_type = client.responses.create(
                model="gpt-5.6-luna",
                input=f"""
Determine the user's MEMORY intent.

Choose exactly one:

SAVE
RETRIEVE

SAVE means:
The user is giving information that NEXORA should remember.

RETRIEVE means:
The user is asking what NEXORA already remembers.

Examples:

"Remember that I am learning Java"
-> SAVE

"Keep in mind that I have an interview Monday"
-> SAVE

"I want you to remember this"
-> SAVE

"What do you remember about me?"
-> RETRIEVE

"Do you know what I am learning?"
-> RETRIEVE

"What have I asked you to remember?"
-> RETRIEVE

USER REQUEST:
{command}

Return ONLY SAVE or RETRIEVE.
"""
            )

            memory_action = (
                memory_type.output_text
                .strip()
                .upper()
            )

            if memory_action not in [
                "SAVE",
                "RETRIEVE"
            ]:

                memory_action = ""

        except Exception as e:

            print(
                "Memory AI classification failed:",
                e
            )

    # -----------------------------------------
    # Local fallback
    # -----------------------------------------

    if memory_action == "":

        text = command.lower()

        retrieve_patterns = [
            "what do you remember",
            "do you remember",
            "what have i asked you to remember",
            "what have i told you",
            "what do you know about me",
            "show my memories",
            "show memories"
        ]

        if any(
            pattern in text
            for pattern in retrieve_patterns
        ):

            memory_action = "RETRIEVE"

        else:

            memory_action = "SAVE"

    # -----------------------------------------
    # SAVE
    # -----------------------------------------

    if memory_action == "SAVE":

        information = command

        if use_ai():

            try:

                memory_response = client.responses.create(
                    model="gpt-5.6-luna",
                    input=f"""
Extract only the important information
the user wants NEXORA to remember.

User request:
{command}

Return only the information to remember.
"""
                )

                extracted_information = (
                    memory_response.output_text
                    .strip()
                )

                if extracted_information:

                    information = extracted_information

            except Exception as e:

                print(
                    "Memory extraction failed:",
                    e
                )

        try:

            tools.save_memory(
                information
            )

        except Exception as e:

            print(
                "Memory save error:",
                e
            )

            return {
                "result": "Unable to save memory."
            }

        return {
            "result": f"Memory saved: {information}"
        }

    # -----------------------------------------
    # RETRIEVE
    # -----------------------------------------

    try:

        memories = tools.get_memories()

    except Exception as e:

        print(
            "Memory retrieval error:",
            e
        )

        return {
            "result": "Unable to retrieve memories."
        }

    if memories:

        return {
            "result": str(memories)
        }

    return {
        "result": "No memories found."
    }


# =========================================================
# RAG NODE
# =========================================================
# =========================================================
# RAG / DOCUMENT NODE
# =========================================================

def rag_node(state: AgentState):

    command = state.get(
        "command",
        ""
    ).strip()

    file_path = state.get(
        "file_path"
    )

    # -----------------------------------------------------
    # Find document
    # -----------------------------------------------------

    if not file_path:

        if os.path.exists("questions.docx"):
            file_path = "questions.docx"

    if not file_path or not os.path.exists(file_path):

        return {
            "result": "Please upload a PDF or document first."
        }

    # -----------------------------------------------------
    # Read document
    # -----------------------------------------------------

    try:

        if file_path.lower().endswith(".pdf"):

            reader = PdfReader(file_path)

            document_text = ""

            for page in reader.pages:

                text = page.extract_text()

                if text:
                    document_text += text + "\n"

        elif file_path.lower().endswith(".docx"):

            document = Document(file_path)

            parts = []

            for paragraph in document.paragraphs:

                if paragraph.text.strip():

                    parts.append(
                        paragraph.text.strip()
                    )

            for table in document.tables:

                for row in table.rows:

                    row_text = " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text.strip()
                    )

                    if row_text:
                        parts.append(row_text)

            document_text = "\n".join(parts)

        elif file_path.lower().endswith(".txt"):

            with open(
                file_path,
                "r",
                encoding="utf-8"
            ) as file:

                document_text = file.read()

        else:

            return {
                "result": "Unsupported document format."
            }

    except Exception as e:

        print(
            "Document reading error:",
            e
        )

        return {
            "result": f"Unable to read the document: {e}"
        }

    if not document_text.strip():

        return {
            "result": "The uploaded document is empty."
        }

    # -----------------------------------------------------
    # Decide document request type using AI
    # -----------------------------------------------------

    request_type = "QUESTION"

    if use_ai():

        try:

            analysis_type = client.responses.create(

                model="gpt-5.6-luna",

                input=f"""
Understand what the user wants from the uploaded document.

Choose exactly one:

ANALYZE
QUESTION

ANALYZE means:
The user wants a summary, overview, key points,
important questions, insights, or a complete analysis
of the document.

Examples:
"Summarize the document"
"Give me the key points"
"What are the important questions from this PDF?"
"Analyze this document"
"Give me summary, key points and important questions"

QUESTION means:
The user wants an answer to a specific question
from the document.

Examples:
"What is PreparedStatement?"
"Explain batch updates"
"What does the document say about ResultSet?"

USER REQUEST:
{command}

Return only ANALYZE or QUESTION.
"""
            )

            value = (
                analysis_type.output_text
                .strip()
                .upper()
            )

            if value in [
                "ANALYZE",
                "QUESTION"
            ]:

                request_type = value

        except Exception as e:

            print(
                "Document request classification failed:",
                e
            )

            text = command.lower()

            if any(word in text for word in [
                "summarize",
                "summary",
                "key points",
                "important questions",
                "analyze",
                "overview"
            ]):

                request_type = "ANALYZE"

    # -----------------------------------------------------
    # DOCUMENT ANALYSIS
    # -----------------------------------------------------

    if request_type == "ANALYZE":

        try:

            response = client.responses.create(

                model="gpt-5.6-luna",

                input=f"""
You are NEXORA's document analysis assistant.

Analyze the uploaded document below.

DOCUMENT:
{document_text}

USER REQUEST:
{command}

Give the result in exactly this structure:

📄 Document Summary

Write a clear and easy-to-understand summary
of the complete document.

🔑 Key Points

Give the most important points from the document.
Use numbered points.

❓ Important Questions

Generate important exam/interview/study questions
that can be answered from the document.
Use numbered questions.

Rules:
- Base everything only on the document.
- Do not invent information.
- Keep the summary concise but useful.
- Focus on the major concepts and topics.
- Important Questions should cover the main topics.
"""
            )

            answer = (
                response.output_text
                .strip()
            )

            if answer:

                return {
                    "result": answer
                }

        except Exception as e:

            print(
                "Document analysis error:",
                e
            )

            return {
                "result": "Unable to analyze the document."
            }

    # -----------------------------------------------------
    # SPECIFIC DOCUMENT QUESTION
    # -----------------------------------------------------

    try:

        response = client.responses.create(

            model="gpt-5.6-luna",

            input=f"""
You are NEXORA.

Answer the user's question using ONLY
the uploaded document below.

DOCUMENT:
{document_text}

USER QUESTION:
{command}

Give a clear, direct and accurate answer.

Do not invent information that is not present
in the document.
"""
        )

        answer = (
            response.output_text
            .strip()
        )

        if answer:

            return {
                "result": answer
            }

    except Exception as e:

        print(
            "Document question error:",
            e
        )

    return {
        "result": "Unable to answer from the document."
    }

# =========================================================
# EMAIL NODE
# =========================================================

def email_node(state: AgentState):

    command = state.get(
        "command",
        ""
    ).strip()

    recipient = ""
    message = ""

    # -----------------------------------------
    # AI extraction
    # -----------------------------------------

    if use_ai():

        try:

            details = client.responses.create(
                model="gpt-5.6-luna",
                input=f"""
Extract the email details from this request:

{command}

Return exactly:

RECIPIENT: email address
MESSAGE: message text
"""
            )

            lines = (
                details.output_text
                .strip()
                .splitlines()
            )

            for line in lines:

                upper_line = line.upper()

                if upper_line.startswith(
                    "RECIPIENT:"
                ):

                    recipient = (
                        line.split(
                            ":",
                            1
                        )[1].strip()
                    )

                elif upper_line.startswith(
                    "MESSAGE:"
                ):

                    message = (
                        line.split(
                            ":",
                            1
                        )[1].strip()
                    )

        except Exception as e:

            print(
                "Email AI extraction failed:",
                e
            )

    # -----------------------------------------
    # Local fallback recipient
    # -----------------------------------------

    if not recipient:

        parts = command.split()

        for part in parts:

            if (
                "@" in part
                and "." in part
            ):

                recipient = part.strip(
                    ".,!?"
                )

                break

    # -----------------------------------------
    # Local fallback message
    # -----------------------------------------

    if not message:

        lower_command = command.lower()

        if "saying" in lower_command:

            message = command.split(
                "saying",
                1
            )[1].strip()

        elif "saying" in command:

            message = command.split(
                "Saying",
                1
            )[1].strip()

        else:

            message = "Hello from NEXORA."

    # -----------------------------------------
    # Validation
    # -----------------------------------------

    if not recipient:

        return {
            "result": "Email address is missing."
        }

    if (
        "@" not in recipient
        or "." not in recipient
    ):

        return {
            "result": "Invalid email address."
        }

    if not message:

        return {
            "result": "Email message is missing."
        }

    print("\nEmail Details:")
    print("Recipient:", recipient)
    print("Message:", message)

    # -----------------------------------------
    # Permission
    # -----------------------------------------

    allowed = ask_permission(
        f"Send an email to {recipient}",
        state
    )

    if not allowed:

        if os.getenv("NEXORA_UI") == "1":

            return {
                "result": "PERMISSION_REQUIRED"
            }

        return {
            "result": "Email cancelled by user."
        }

    print("Permission granted.")

    try:

        result = tools.send_email(
            recipient,
            message,
            state.get(
                "file_path"
            ),
            state.get(
                "google_access_token"
            )
        )

        print(
            "\n" + result
        )

        return {
            "result": result
        }

    except Exception as e:

        print(
            "Email error:",
            e
        )

        return {
            "result": f"Email could not be sent: {e}"
        }


# =========================================================
# OTHER NODE
# =========================================================

def other_node(state: AgentState):

    return {
        "result": "Unknown request. Please try another command."
    }


# =========================================================
# ROUTING
# =========================================================

def route_request(state: AgentState):

    action = state.get(
        "action",
        "OTHER"
    )

    if action == "MEETING":

        return "meeting"

    elif action == "ADD_TASK":

        return "task"

    elif action == "SHOW_TASK":

        return "show_task"

    elif action == "SHOW_MEETING":

        return "show_meeting"

    elif action == "MEMORY":

        return "memory"

    elif action == "RAG":

        return "rag"

    elif action == "SEND_EMAIL":

        return "email"

    return "other"


# =========================================================
# LANGGRAPH
# =========================================================

graph = StateGraph(
    AgentState
)


graph.add_node(
    "analyze",
    analyze_request
)


graph.add_node(
    "meeting",
    meeting_node
)


graph.add_node(
    "task",
    task_node
)


graph.add_node(
    "show_task",
    show_task_node
)


graph.add_node(
    "show_meeting",
    show_meeting_node
)


graph.add_node(
    "memory",
    memory_node
)


graph.add_node(
    "rag",
    rag_node
)


graph.add_node(
    "email",
    email_node
)


graph.add_node(
    "other",
    other_node
)


graph.add_edge(
    START,
    "analyze"
)


graph.add_conditional_edges(
    "analyze",
    route_request,
    {
        "meeting": "meeting",
        "task": "task",
        "show_task": "show_task",
        "show_meeting": "show_meeting",
        "memory": "memory",
        "rag": "rag",
        "email": "email",
        "other": "other"
    }
)


graph.add_edge(
    "meeting",
    END
)


graph.add_edge(
    "task",
    END
)


graph.add_edge(
    "show_task",
    END
)


graph.add_edge(
    "show_meeting",
    END
)


graph.add_edge(
    "memory",
    END
)


graph.add_edge(
    "rag",
    END
)


graph.add_edge(
    "email",
    END
)


graph.add_edge(
    "other",
    END
)


app = graph.compile()


# =========================================================
# TEST
# =========================================================

if __name__ == "__main__":

    result = app.invoke({

        "command": "I want to study Java tomorrow",

        "action": "",

        "result": "",

        "file_path": None,

        "permission_granted": False,

        "google_access_token": None
    })

    print("\nFinal Result:")

    print(
        result.get(
            "result",
            "No result."
        )
    )