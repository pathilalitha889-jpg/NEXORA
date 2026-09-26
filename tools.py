import os
import sqlite3
import base64
import mimetypes

from email.message import EmailMessage
from datetime import datetime, timedelta

from dotenv import load_dotenv
from docx import Document

from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build


# ==============================
# CONFIGURATION
# ==============================

load_dotenv()

DATABASE = "nexora.db"

GOOGLE_SCOPES = [
    "https://www.googleapis.com/auth/gmail.send",
    "https://www.googleapis.com/auth/calendar.events"
]


# ==============================
# DATABASE SETUP
# ==============================

def get_connection():

    return sqlite3.connect(DATABASE)


conn = get_connection()
cursor = conn.cursor()

cursor.execute("""
CREATE TABLE IF NOT EXISTS tasks (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    task TEXT
)
""")

cursor.execute("""
CREATE TABLE IF NOT EXISTS meetings (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    person TEXT,
    date TEXT,
    time TEXT
)
""")

cursor.execute("""
CREATE TABLE IF NOT EXISTS memories (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    information TEXT
)
""")

conn.commit()
conn.close()


# ==============================
# TASKS
# ==============================

def add_task(task):

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        "INSERT INTO tasks (task) VALUES (?)",
        (task,)
    )

    conn.commit()
    conn.close()

    return "Task Added!"


def get_tasks():

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        "SELECT * FROM tasks"
    )

    tasks = cursor.fetchall()

    conn.close()

    return tasks


# ==============================
# MEETINGS
# ==============================

def save_meeting(person, date, time):

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        """
        INSERT INTO meetings (person, date, time)
        VALUES (?, ?, ?)
        """,
        (person, date, time)
    )

    conn.commit()
    conn.close()

    return "Meeting saved!"


def get_meetings():

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        "SELECT * FROM meetings"
    )

    meetings = cursor.fetchall()

    conn.close()

    return meetings


# ==============================
# GOOGLE AUTHENTICATION
# ==============================

def get_google_credentials(access_token):

    if not access_token:
        raise ValueError(
            "Google access token is missing. Please login with Google again."
        )

    credentials = Credentials(
        token=access_token,
        scopes=GOOGLE_SCOPES
    )

    return credentials


# ==============================
# GOOGLE CALENDAR
# ==============================

def create_calendar_event(
    person,
    date,
    time,
    access_token=None
):

    credentials = get_google_credentials(
        access_token
    )

    service = build(
        "calendar",
        "v3",
        credentials=credentials
    )

    start_datetime = datetime.strptime(
        f"{date} {time}",
        "%Y-%m-%d %H:%M"
    )

    end_datetime = (
        start_datetime
        + timedelta(minutes=30)
    )

    timezone = "Asia/Kolkata"

    start_time = (
        start_datetime.isoformat()
        + "+05:30"
    )

    end_time = (
        end_datetime.isoformat()
        + "+05:30"
    )

    event = {

        "summary": f"Meeting with {person}",

        "description": (
            "Meeting scheduled by NEXORA "
            "Intelligent Multi-Tool AI Agent."
        ),

        "start": {
            "dateTime": start_time,
            "timeZone": timezone
        },

        "end": {
            "dateTime": end_time,
            "timeZone": timezone
        },

        "reminders": {
            "useDefault": False,
            "overrides": [
                {
                    "method": "popup",
                    "minutes": 10
                },
                {
                    "method": "email",
                    "minutes": 1440
                }
            ]
        }
    }

    created_event = service.events().insert(
        calendarId="primary",
        body=event
    ).execute()

    return created_event.get(
        "htmlLink",
        "Calendar event created successfully!"
    )


# Keep compatibility with older code
def schedule_meeting(
    person,
    date,
    time,
    access_token=None
):

    return create_calendar_event(
        person,
        date,
        time,
        access_token
    )


# ==============================
# GMAIL
# ==============================

def send_email(
    recipient,
    message,
    attachment_path=None,
    access_token=None
):

    credentials = get_google_credentials(
        access_token
    )

    service = build(
        "gmail",
        "v1",
        credentials=credentials
    )

    email = EmailMessage()

    email["To"] = recipient
    email["Subject"] = "Message from NEXORA"

    email.set_content(message)

    # Optional attachment
    if attachment_path:

        if os.path.exists(attachment_path):

            mime_type, _ = mimetypes.guess_type(
                attachment_path
            )

            if mime_type:

                main_type, sub_type = (
                    mime_type.split("/", 1)
                )

            else:

                main_type = "application"
                sub_type = "octet-stream"

            with open(
                attachment_path,
                "rb"
            ) as file:

                file_data = file.read()

            filename = os.path.basename(
                attachment_path
            )

            email.add_attachment(
                file_data,
                maintype=main_type,
                subtype=sub_type,
                filename=filename
            )

    encoded_message = base64.urlsafe_b64encode(
        email.as_bytes()
    ).decode()

    body = {
        "raw": encoded_message
    }

    sent_message = (
        service.users()
        .messages()
        .send(
            userId="me",
            body=body
        )
        .execute()
    )

    if sent_message.get("id"):

        return "Email sent successfully!"

    return "Unable to send email."


# ==============================
# MEMORY
# ==============================

def save_memory(information):

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        """
        INSERT INTO memories (information)
        VALUES (?)
        """,
        (information,)
    )

    conn.commit()
    conn.close()

    return "Memory saved!"


def get_memories():

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        "SELECT * FROM memories"
    )

    memories = cursor.fetchall()

    conn.close()

    return memories


# ==============================
# RAG DOCUMENT
# ==============================

def read_document():

    document = Document(
        "questions.docx"
    )

    text = ""

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text += (
                paragraph.text
                + "\n"
            )

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():

                    text += (
                        cell.text
                        + "\n"
                    )

    return text


def search_document(question):

    document_text = read_document()

    question_words = (
        question.lower()
        .replace("?", "")
        .replace(".", "")
        .replace(",", "")
        .split()
    )

    words = set(
        question_words
    )

    best_score = 0
    best_text = ""

    paragraphs = (
        document_text.split("\n")
    )

    for paragraph in paragraphs:

        paragraph_lower = (
            paragraph.lower()
        )

        score = 0

        for word in words:

            if (
                len(word) > 2
                and word in paragraph_lower
            ):

                score += 1

        if score > best_score:

            best_score = score
            best_text = paragraph

    if best_score == 0:

        return document_text

    return best_text